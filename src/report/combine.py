"""Combined multi-locale workbook: merges audited workbooks (one per locale) into one.

Tabs: All Locales (overview, tier totals with the reference rates, one row per locale), one tab per locale holding that
locale's AI Review (trigger sentences in bold red and trigger words underlined, as in the source), then Trigger Words,
Agents (all) and Annotated Master with a Locale column. It reads the saved audit workbooks only; nothing is re-scored."""
import logging
import os
import re
from collections import Counter
from typing import Dict, List, Optional, Sequence

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import excel_writer as ew
from .tiers import HIGH, HUMAN, LIKELY_AI, NOT_CHECKED, TIER_ORDER

logger = logging.getLogger(__name__)

TIERS = tuple(t for t in TIER_ORDER if t != HUMAN)  # rater rationales are never called Likely human
INFO_KEYS = ("Model", "Columns checked", "Average justification quality", "Trigger sentences and words",
             "Structural errors", "Structural warnings")
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


class CombineError(ValueError):
    """The audited workbooks cannot be combined (too few, or not audit workbooks)."""


def _tier_of(text) -> str:
    text = str(text or "").strip()
    if not text or text.startswith(NOT_CHECKED):
        return NOT_CHECKED
    return HIGH if text.startswith(HIGH.split(" –")[0]) else text


def _verdict_col(master: List[list]) -> Optional[int]:
    """Where the QA Verdict column sits in a locale's Annotated Master; it follows that locale's own input columns."""
    header = [str(h or "").strip() for h in master[0]] if master else []
    return header.index("QA Verdict") if "QA Verdict" in header else None


def _number(value):
    """Cells read back as text become numbers again, so the combined sheets sort and chart."""
    if isinstance(value, str) and re.fullmatch(r"-?\d+(\.\d+)?", value.strip()):
        number = float(value)
        return int(number) if number.is_integer() and "." not in value else number
    return value


def _cell_value(value):
    if isinstance(value, str):
        return value[:ew.EXCEL_CELL_LIMIT]
    return value  # numbers, and CellRichText (the red and underlined runs), pass through


def _sheet_name(name: str, taken: set) -> str:
    base = _BAD_SHEET_CHARS.sub("-", name).strip()[:31] or "Locale"
    candidate, n = base, 2
    while candidate.lower() in taken:
        candidate = f"{base[:28]}-{n}"
        n += 1
    taken.add(candidate.lower())
    return candidate


def _rows(ws) -> List[list]:
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _read(path: str, fallback_name: str) -> Dict:
    wb = openpyxl.load_workbook(path, data_only=True, rich_text=True)
    missing = [s for s in ("Summary Dashboard", "AI Review", "Annotated Master", "Agents") if s not in wb.sheetnames]
    if missing:
        raise CombineError(f"{os.path.basename(path)} is not an audit workbook (missing: {', '.join(missing)}).")
    info, meta, notes, in_tiers = {}, {}, [], False
    for row in wb["Summary Dashboard"].iter_rows(values_only=True):
        key = row[0]
        if key == "Tier":
            in_tiers = True
        elif in_tiers and _tier_of(key) in TIERS and key is not None:
            meta[_tier_of(key)] = list(row[2:7])
        elif in_tiers and key and row[1] is None:
            notes.append(key)
        elif key in INFO_KEYS + ("Locale",) and row[1] is not None:
            info[key] = row[1]
    master = _rows(wb["Annotated Master"])
    return {"name": str(info.get("Locale") or fallback_name), "info": info, "meta": meta, "notes": notes,
            "agents": _rows(wb["Agents"]), "words": _rows(wb["Trigger Words"]) if "Trigger Words" in wb.sheetnames else [],
            "review": _rows(wb["AI Review"]), "master": master}


def _table(ws, top: int, headers: Sequence[str], rows: Sequence[Sequence], formats: Optional[Dict[int, str]] = None) -> int:
    ew._header(ws, top, list(headers))
    for i, row in enumerate(rows, top + 1):
        for j, value in enumerate(row, 1):
            cell = ws.cell(i, j, value)
            cell.font, cell.border = ew.FONT_DEFAULT, ew.THIN_BORDER
            if formats and j in formats and isinstance(value, (int, float)):
                cell.number_format = formats[j]
    return top + len(rows) + 1


def _tier_fill(ws, row: int, col: int) -> None:
    fill = ew.TIER_FILLS.get(_tier_of(ws.cell(row, col).value))
    if fill:
        ws.cell(row, col).fill = fill


def _overview(wb, locales: List[Dict]) -> None:
    ws = wb.create_sheet("All Locales")
    ws.sheet_view.showGridLines = False
    models = sorted({str(d["info"].get("Model")) for d in locales if d["info"].get("Model")})
    ws["A1"], ws["A1"].font = "COMBINED EVALUATION QA AUDIT REPORT", Font(name="Calibri", size=16, bold=True)
    ws["A2"] = f"{len(locales)} locales · model {', '.join(models)} · each locale tab holds that locale's AI Review"
    ws["A2"].font = ew.FONT_NOTE
    total = Counter()
    for d in locales:
        verdict = _verdict_col(d["master"])
        d["counts"] = Counter() if verdict is None else Counter(
            _tier_of(r[verdict]) for r in d["master"][1:] if r[0] is not None and len(r) > verdict)
        d["n"] = sum(d["counts"].values())
        total.update(d["counts"])
        d["agent_n"] = sum(1 for r in d["agents"][1:] if r[0])
        d["agent_likely"] = sum(1 for r in d["agents"][1:] if r[0] and (_number(r[2]) or 0) > 0)
    n_all = sum(total.values())
    overview = (("Locales", len(locales)), ("Rows", n_all), ("Agents (locale-agent pairs)", sum(d["agent_n"] for d in locales)),
                ("Agents with any Likely AI", f"{sum(d['agent_likely'] for d in locales)} of {sum(d['agent_n'] for d in locales)}"),
                ("Structural errors", sum(_number(d["info"].get("Structural errors")) or 0 for d in locales)))
    ws["A4"], ws["A4"].font = "Overview", ew.FONT_BOLD
    for i, (key, value) in enumerate(overview, 5):
        ws.cell(i, 1, key).font = ew.FONT_BOLD
        ws.cell(i, 2, value).font, ws.cell(i, 2).alignment = ew.FONT_DEFAULT, Alignment(horizontal="left")
    top = 11
    ws.cell(top - 1, 1, "Verdict tiers, all locales").font = ew.FONT_BOLD
    meta = next((d["meta"] for d in locales if d["meta"]), {})
    rows = [[t, total.get(t, 0), total.get(t, 0) / n_all if n_all else 0] + (meta.get(t) or [None] * 5) for t in TIERS]
    pct = {3: "0.0%", 5: "0.0%", 6: "0.0%", 7: "0.0%", 8: "0.0%"}
    end = _table(ws, top, ["Tier", "Rows", "Share of rows", "Meaning"] + [c[2] for c in ew.RATE_COLUMNS], rows, pct)
    for r in range(top + 1, end):
        _tier_fill(ws, r, 1)
    top = end + 2
    ws.cell(top - 1, 1, "By locale (each locale tab holds that locale's AI Review)").font = ew.FONT_BOLD
    headers = ["Locale", "Rows", "Agents", "Agents with Likely AI"] + list(TIERS) + [
        "% Likely AI", "% High score", "% Not checked", "Avg justification quality", "Structural errors", "Trigger search"]
    k = len(TIERS)
    body = []
    for d in locales:
        c, n = d["counts"], d["n"] or 1
        body.append([d["sheet"], d["n"], d["agent_n"], d["agent_likely"]] + [c.get(t, 0) for t in TIERS] +
                    [c.get(LIKELY_AI, 0) / n, c.get(HIGH, 0) / n, c.get(NOT_CHECKED, 0) / n,
                     _number(d["info"].get("Average justification quality")), _number(d["info"].get("Structural errors")),
                     d["info"].get("Trigger sentences and words")])
    end = _table(ws, top, headers, body, {5 + k: "0.0%", 6 + k: "0.0%", 7 + k: "0.0%"})
    for r, d in enumerate(locales, top + 1):
        ws.cell(r, 1).hyperlink = f"#'{d['sheet']}'!A1"
        ws.cell(r, 1).font = Font(name="Calibri", size=10, color="0563C1", underline="single")
    ws.cell(end, 1, "Total").font = ew.FONT_BOLD
    col = lambda j: get_column_letter(j)
    for j in range(2, 5 + k):
        ws.cell(end, j, f"=SUM({col(j)}{top + 1}:{col(j)}{end - 1})").font = ew.FONT_BOLD
    for j, tier_col in ((5 + k, 5), (6 + k, 6), (7 + k, 4 + k)):
        ws.cell(end, j, f"={col(tier_col)}{end}/B{end}").number_format = "0.0%"
    note_row = end + 2
    for note in next((d["notes"] for d in locales if d["notes"]), ()):
        ws.cell(note_row, 1, note).font = ew.FONT_NOTE
        note_row += 1
    ew._widths(ws, {1: 32, 2: 10, 3: 14, 4: 40, **{j: 14 for j in range(5, 8 + k + 3)}, 8 + k + 2: 30})
    ws.freeze_panes = "A4"


FIX_HEADER = "Rationale Fix (plain-language rewrite)"


def _review_tab(wb, d: Dict, fixes: Optional[Dict[str, str]] = None) -> None:
    """One locale's AI Review. With `fixes` ({"<tab>|<row>": text}) a labeled rewrite column sits right after Rationale."""
    ws = wb.create_sheet(d["sheet"])
    headers = [str(h) for h in d["review"][0]]
    at = headers.index("Rationale") + 1 if fixes and "Rationale" in headers else None  # position of the new column
    if at is not None:
        headers.insert(at, FIX_HEADER)
    ew._header(ws, 1, headers)
    r = 2
    for row in d["review"][1:]:
        if all(v is None for v in row):
            continue
        if at is not None:
            row = list(row)
            row.insert(at, (fixes or {}).get(f"{d['sheet']}|{r}"))
        for j, value in enumerate(row, 1):
            cell = ws.cell(r, j, value if not isinstance(value, str) or len(value) > 40 else _number(value))
            if isinstance(value, str):
                cell.value = _cell_value(cell.value)
            cell.font = ew.FONT_DEFAULT
            if headers[j - 1] in ("Rationale", FIX_HEADER, "Trigger", "Each Column Alone"):
                cell.alignment = ew.WRAP
        _tier_fill(ws, r, headers.index("Verdict") + 1)
        r += 1
    by_name = {"Row #": 8, "Task ID": 20, "Agent": 30, "Verdict": 22, "AI Score": 10, "Bar": 10, "Decided On": 22,
               "Rationale": 60, FIX_HEADER: 60, "Trigger": 50, "Each Column Alone": 40, "Human Reviewer Decision": 22,
               "Reviewer Notes": 30}
    widths = {j: by_name.get(h, 14) for j, h in enumerate(headers, 1)}
    ew._widths(ws, widths)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(r - 1, 1)}"
    if "Human Reviewer Decision" in headers:
        letter = get_column_letter(headers.index("Human Reviewer Decision") + 1)
        validation = DataValidation(type="list", formula1=ew.REVIEW_CHOICES, allow_blank=True)
        ws.add_data_validation(validation)
        validation.add(f"{letter}2:{letter}{max(r - 1, 2)}")


def _stacked(wb, name: str, key: str, locales: List[Dict], widths: Dict[int, float], wrap=()) -> None:
    sources = [d for d in locales if d[key]]
    if not sources:
        return
    ws = wb.create_sheet(name)
    headers = ["Locale"] + [str(h) for h in sources[0][key][0]]
    ew._header(ws, 1, headers)
    r = 2
    for d in sources:
        for row in d[key][1:]:
            if all(v is None for v in row):
                continue
            ws.cell(r, 1, d["sheet"]).font = ew.FONT_DEFAULT
            for j, value in enumerate(row, 2):
                cell = ws.cell(r, j, _cell_value(value if not isinstance(value, str) or len(value) > 40 else _number(value)))
                cell.font = ew.FONT_DEFAULT
                if j - 2 in wrap:
                    cell.alignment = ew.WRAP
            r += 1
    ew._widths(ws, widths)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(r - 1, 1)}"


def combine_workbooks(paths: Sequence[str], output_path: str, fixes: Optional[Dict[str, str]] = None) -> Dict:
    """Merges audited workbooks (one per locale, as written by the audit) into one combined workbook."""
    paths = [p for p in paths if p]
    if len(paths) < 2:
        raise CombineError("Combining needs at least two audited workbooks.")
    taken: set = {"all locales", "trigger words", "agents (all)", "annotated master"}
    locales = []
    for path in paths:
        stem = os.path.splitext(os.path.basename(path))[0]
        data = _read(path, re.sub(r"^audited_[0-9a-f]+_", "", stem))
        data["sheet"] = _sheet_name(data["name"], taken)
        locales.append(data)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _overview(wb, locales)
    for d in locales:
        _review_tab(wb, d, fixes)
    _stacked(wb, "Trigger Words", "words", locales, {1: 10, 2: 36, 3: 12, 4: 10, 5: 14, 6: 14, 7: 12, 8: 50})
    _stacked(wb, "Agents (all)", "agents", locales, {1: 10, 2: 36, 3: 8, 4: 10, 5: 16, 6: 12, 7: 14, 8: 14, 9: 12})
    _stacked(wb, "Annotated Master", "master", locales,
             {1: 10, 2: 38, 3: 30, 4: 10, 5: 10, 6: 30, 7: 50, 8: 24, 9: 10, 10: 10, 11: 14, 12: 50, 13: 40, 14: 10, 15: 26, 16: 10},
             wrap=(4, 5, 10, 11))
    wb.save(output_path)
    rows = sum(d["n"] for d in locales)
    logger.info("combined %d audited workbooks (%d rows) into %s", len(locales), rows, output_path)
    return {"locales": [d["sheet"] for d in locales], "rows": rows, "output": output_path}
