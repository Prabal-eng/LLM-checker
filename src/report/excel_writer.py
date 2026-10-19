"""Audit workbook (v2, spec 2026-10-01): Summary Dashboard, Agents, Trigger Words (when triggers ran), AI Review,
Annotated Master, Technical, and Structural Issues when there are any. One verdict tier per row; trigger sentences
in bold red, trigger words underlined (2026-10-02)."""
import logging
from typing import Any, Dict, List, Optional

import openpyxl
import pandas as pd
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..qa.quality import QualityAssessment
from ..qa.structural import StructuralIssue
from .rows import RowReport
from .tier_rates import RATED_TIERS
from .tiers import EXPLAINED, HIGH, HUMAN, LIKELY_AI, MODERATE, NOT_CHECKED, SECOND_LOOK, TIER_ORDER, TierRules, bar_for
from .triggers import JOIN, TriggerResult, pct

logger = logging.getLogger(__name__)

EXCEL_CELL_LIMIT = 32_767


def _fill(rgb: str) -> PatternFill:
    return PatternFill(start_color=rgb, end_color=rgb, fill_type="solid")


NAVY_HEADER, GRAY_HEADER = _fill("1F4E78"), _fill("595959")
RED_FILL, YELLOW_FILL = _fill("FFC7CE"), _fill("FFEB9C")
TIER_FILLS = {LIKELY_AI: _fill("F4A7A7"), HIGH: _fill("FFC7CE"), MODERATE: _fill("FFEB9C"),
              SECOND_LOOK: _fill("FFF2CC"), NOT_CHECKED: _fill("D9D9D9"), HUMAN: _fill("C6EFCE")}
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_BOLD = Font(name="Calibri", size=11, bold=True)
FONT_DEFAULT = Font(name="Calibri", size=10)
FONT_NOTE = Font(name="Calibri", size=10, italic=True)
TRIGGER_FONT = InlineFont(b=True, color="C00000")             # a trigger sentence
WORD_FONT = InlineFont(b=True, u="single", color="C65911")     # a trigger word
WORD_IN_TRIGGER_FONT = InlineFont(b=True, u="single", color="C00000")
SENTENCE, WORD = "sentence", "word"
LABEL_FONT = InlineFont(b=True)
THIN_BORDER = Border(*(Side(style="thin", color="D9D9D9"),) * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
REVIEW_CHOICES = '"Approved (Human),Confirmed AI,False Positive,Needs Discussion"'
RATE_COLUMNS = (("25-100", "human", "Human texts, 25-100 words"), ("25-100", "ai", "AI texts, 25-100 words"),
                ("all", "human", "Human texts, all lengths"), ("all", "ai", "AI texts, all lengths"))
NOTES = (
    "Likely AI is the only decision. Its bar depends on the text's length (see the Bar column in AI Review) and "
    "is set so that under 1% of human writing in validation clears it.",
    "The other tiers rank rows for review by AI score; they are not decisions.",
    "Rater rationales are never called Likely human: missed AI on rater text has not been measured, so a low "
    "score is not proof of human writing.",
    "Trigger sentences and words are indicative: the model scores the whole text. The sentence in bold red, or the "
    "underlined word, is what takes away the largest share of the row's AI signal when removed. The AI signal is "
    "the score's log-odds: near 99% a large share barely moves the percent (99% -> 97% is a quarter of it).",
    "Trigger Words: a word or phrase that triggers in many rows and appears in most rows is likely shared wording "
    "(such as rubric terms) that the model over-reacts to. That points at the checker, not at the agents.",
)


def verdict_text(report: RowReport) -> str:
    return f"{report.tier} ({report.reason})" if report.reason else report.tier


def trigger_text(report: RowReport) -> str:
    if report.trigger is None:
        return "not run" if report.tier in EXPLAINED else ""
    return report.trigger.summary()


def trigger_sentences(report: RowReport) -> str:
    if report.trigger is not None and report.trigger.kind == "sentences":
        return " | ".join(s.text for s, _ in report.trigger.sentences)
    if report.trigger is not None and report.trigger.words:
        return ""  # the words are in their own column
    return trigger_text(report)


def trigger_words(report: RowReport) -> str:
    if report.trigger is None:
        return ""
    return "; ".join(f"{w.text} ({pct(w.share)})" for w in report.trigger.words)


def drops_text(trigger: Optional[TriggerResult]) -> str:
    if trigger is None or not trigger.drops:
        return ""
    return "\n".join(f"{-share:+.0%} of signal: [{s.column}] {s.text[:80]}" for s, share in trigger.drops)


def word_drops_text(trigger: Optional[TriggerResult]) -> str:
    if trigger is None or not trigger.word_drops:
        return ""
    return "\n".join(f'{-w.share:+.0%} of signal: "{w.text}" [{w.column}]' for w in trigger.word_drops)


def _font(kinds: set) -> Optional[InlineFont]:
    if WORD in kinds:
        return WORD_IN_TRIGGER_FONT if SENTENCE in kinds else WORD_FONT
    return TRIGGER_FONT if SENTENCE in kinds else None


def rationale_cell(report: RowReport):
    """The decided text, each cell labelled when there are several, with trigger sentences in bold red and trigger
    words underlined (red inside a trigger sentence, orange elsewhere)."""
    marks: Dict[str, list] = {}
    trigger = report.trigger
    if trigger is not None:
        if trigger.kind == "sentences":
            for sentence, _ in trigger.sentences:
                marks.setdefault(sentence.column, []).append((sentence.start, sentence.end, SENTENCE))
        for word in trigger.words:
            marks.setdefault(word.column, []).append((word.start, word.end, WORD))
    parts: list = []
    for i, (column, text) in enumerate(report.decided):
        if i:
            parts.append(JOIN)
        if len(report.decided) > 1:
            parts.append(TextBlock(LABEL_FONT, f"[{column}] "))
        spans = [(a, b, k) for a, b, k in marks.get(column, []) if 0 <= a < b <= len(text)]
        cuts = sorted({0, len(text)} | {p for a, b, _ in spans for p in (a, b)})
        for a, b in zip(cuts, cuts[1:]):
            font = _font({k for s, e, k in spans if s <= a and b <= e})
            parts.append(TextBlock(font, text[a:b]) if font else text[a:b])
    plain = "".join(p if isinstance(p, str) else p.text for p in parts)
    if len(plain) > EXCEL_CELL_LIMIT or not any(isinstance(p, TextBlock) for p in parts):
        return plain[:EXCEL_CELL_LIMIT]
    return CellRichText(*_without_blank_runs(parts))


def _without_blank_runs(parts: list) -> list:
    """Each whitespace-only piece joins its neighbour (the one before it when there is one): openpyxl writes a
    whitespace-only run without xml:space="preserve", and Excel repairs that on opening as broken "String
    properties" (the space between two adjacent trigger words, or the blank line between two cells)."""
    out: list = []
    lead = ""
    for p in parts:
        text = p if isinstance(p, str) else p.text
        if text.isspace() and out:
            last = out[-1]
            out[-1] = last + text if isinstance(last, str) else TextBlock(last.font, last.text + text)
        elif text.isspace():
            lead += text
        elif text:
            out.append(lead + p if isinstance(p, str) else TextBlock(p.font, lead + p.text))
            lead = ""
    return out


def column_scores(report: RowReport, rules: TierRules) -> str:
    results = report.result.column_results
    if len(results) < 2:
        return ""
    lines = []
    for name, res in results.items():
        tier, reason = rules.tier(res)
        detail = f"{res.score:.0%}" if res.score is not None else reason
        lines.append(f"{name}: {tier} ({detail})")
    return "\n".join(lines)


def _header(ws, row: int, headers: List[str], fill: PatternFill = NAVY_HEADER) -> None:
    for col, name in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=name)
        cell.font, cell.fill = FONT_HEADER, fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws, widths: Dict[int, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _autofit_columns(ws, max_len_cap: int = 50) -> None:
    for col in ws.columns:
        longest = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len_cap, max(longest + 3, 12))


def _box(ws, row: int, ncols: int, wrap: bool = False) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.border = THIN_BORDER
        if wrap:
            cell.alignment = WRAP


class AuditedWorkbookBuilder:
    """Builds the audit workbook from the pipeline's row reports."""

    def __init__(self, tier_rules: TierRules = TierRules()):
        self.tier_rules = tier_rules
        self.wb = openpyxl.Workbook()
        self.wb.remove(self.wb.active)

    def add_summary_sheet(self, summary: Dict[str, Any], rates: Optional[dict], warning: Optional[str] = None):
        ws = self.wb.create_sheet(title="Summary Dashboard")
        ws.cell(row=1, column=1, value="EVALUATION QA AUDIT REPORT").font = Font(
            name="Calibri", size=16, bold=True, color="1F4E78")
        ws.cell(row=2, column=1, value="AI authorship tiers, trigger sentences and structural checks").font = FONT_NOTE
        if warning:  # e.g. an unvalidated preview: shown above every figure
            banner = ws.cell(row=3, column=1, value=warning)
            banner.font = Font(name="Calibri", size=11, bold=True, color="9C0006")
            banner.fill = RED_FILL
        agents = summary.get("agents") or []
        triggers = summary.get("triggers") or {}
        facts = [
            ("Model", summary.get("detector_version")),
            ("Locale", summary.get("locale") or "all rows"),
            ("Sheet", summary.get("detected_sheet")),
            ("Agent column", summary.get("agent_column") or "none found"),
            ("Columns checked", ", ".join(summary.get("text_columns") or [])),
            ("Rows checked", summary.get("total_rows")),
            ("Agents with any Likely AI", f"{sum(1 for a in agents if a['likely_ai'])} of {len(agents)}"),
            ("Average justification quality", f"{(summary.get('avg_quality_score') or 0) * 100:.1f}%"),
            ("Trigger sentences and words",
             f"looked for in {triggers.get('rows', 0)} rows ({triggers.get('seconds', 0):.0f} s)"
             if triggers.get("enabled") else "off"),
        ]
        if summary.get("structural_errors_count") or summary.get("structural_warnings_count"):
            facts += [("Structural errors", summary.get("structural_errors_count", 0)),
                      ("Structural warnings", summary.get("structural_warnings_count", 0))]
        row = 5
        _header(ws, row, ["Run", "Value"])
        for label, value in facts:
            row += 1
            ws.cell(row=row, column=1, value=label).font = FONT_DEFAULT
            ws.cell(row=row, column=2, value=value).font = FONT_BOLD
            _box(ws, row, 2)

        row += 2
        ws.cell(row=row, column=1, value="Verdict tiers (strongest first)").font = FONT_BOLD
        row += 1
        _header(ws, row, ["Tier", "Rows", "Meaning"] + [label for _, _, label in RATE_COLUMNS], GRAY_HEADER)
        meanings, counts = self.tier_rules.meanings(), summary.get("tier_counts") or {}
        for tier in TIER_ORDER:
            if tier == HUMAN and not counts.get(tier):
                continue  # rater rationales are flag-only, so this row would always read 0
            row += 1
            tier_cell = ws.cell(row=row, column=1, value=tier)
            if tier in TIER_FILLS:
                tier_cell.fill = TIER_FILLS[tier]
            ws.cell(row=row, column=2, value=counts.get(tier, 0)).font = FONT_BOLD
            ws.cell(row=row, column=3, value=meanings[tier]).font = FONT_DEFAULT
            if tier in RATED_TIERS:
                if rates:
                    for col, (slice_name, who, _) in enumerate(RATE_COLUMNS, 4):
                        cell = ws.cell(row=row, column=col, value=rates["slices"][slice_name]["tiers"][tier][who])
                        cell.number_format = "0.0%"
                else:
                    ws.cell(row=row, column=4, value="not measured for this model").font = FONT_NOTE
            _box(ws, row, 3 + len(RATE_COLUMNS), wrap=True)

        row += 2
        notes = list(NOTES)
        if rates:
            every = rates["slices"]["all"]
            notes.append(f"Rates: the share of {every['human_texts']:,} human and {every['ai_texts']:,} AI validation "
                         f"texts that land in each tier (measured {rates['measured_at'][:10]}).")
            notes += rates.get("notes", [])
        for note in notes:
            ws.cell(row=row, column=1, value=note).font = FONT_NOTE
            row += 1
        _widths(ws, {1: 34, 2: 40, 3: 60, 4: 15, 5: 15, 6: 15, 7: 15})

    def add_agents_sheet(self, agents: List[dict]):
        ws = self.wb.create_sheet(title="Agents")
        headers = ["Agent", "Rows", LIKELY_AI, HIGH, MODERATE, NOT_CHECKED, "Highest AI Score", "Share Likely AI"]
        _header(ws, 1, headers)
        for row, a in enumerate(agents, 2):
            values = [a["agent"], a["rows"], a["likely_ai"], a["high"], a["moderate"], a["not_checked"],
                      a["highest_score"], a["share_likely_ai"]]
            for col, value in enumerate(values, 1):
                ws.cell(row=row, column=col, value=value).font = FONT_DEFAULT
            ws.cell(row=row, column=7).number_format = ws.cell(row=row, column=8).number_format = "0.0%"
            if a["likely_ai"]:
                ws.cell(row=row, column=3).fill = TIER_FILLS[LIKELY_AI]
            if a["high"]:
                ws.cell(row=row, column=4).fill = TIER_FILLS[HIGH]
            _box(ws, row, len(headers))
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(agents) + 1}"
        _widths(ws, {1: 34, 2: 8, 3: 11, 4: 22, 5: 18, 6: 12, 7: 16, 8: 15})

    def add_trigger_words_sheet(self, table: List[dict]):
        """One line per trigger word or phrase across the sheet, with how many of all rows contain it."""
        ws = self.wb.create_sheet(title="Trigger Words")
        headers = ["Word or Phrase", "Rows It Triggered", "Agents", "Average Share of AI Signal", "Rows Containing It",
                   "Share of All Rows", "Agent Names"]
        _header(ws, 1, headers)
        if not table:
            ws.cell(row=2, column=1, value="No trigger words found").font = FONT_NOTE
        for row, e in enumerate(table, 2):
            names = ", ".join(e["agent_names"][:10]) + (f" and {len(e['agent_names']) - 10} more"
                                                          if len(e["agent_names"]) > 10 else "")
            values = [e["phrase"], e["trigger_rows"], e["agents"], e["average_share"], e["rows_containing"],
                      e["share_of_rows"], names]
            for col, value in enumerate(values, 1):
                ws.cell(row=row, column=col, value=value).font = FONT_DEFAULT
            ws.cell(row=row, column=4).number_format = ws.cell(row=row, column=6).number_format = "0%"
            _box(ws, row, len(headers), wrap=True)
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(table), 1) + 1}"
        _widths(ws, {1: 30, 2: 12, 3: 9, 4: 14, 5: 12, 6: 11, 7: 60})

    def add_ai_review_sheet(self, reports: List[RowReport]):
        """One row per task, strongest first, with the full decided text and its trigger sentences."""
        ws = self.wb.create_sheet(title="AI Review")
        with_ids = any(r.task_id for r in reports)
        headers = (["Row #"] + (["Task ID"] if with_ids else []) +
                   ["Agent", "Verdict", "AI Score", "Bar", "Decided On", "Rationale", "Trigger", "Each Column Alone",
                    "Human Reviewer Decision", "Reviewer Notes"])
        _header(ws, 1, headers)
        col = {name: i for i, name in enumerate(headers, 1)}
        choices = DataValidation(type="list", formula1=REVIEW_CHOICES, allow_blank=True)
        ws.add_data_validation(choices)
        order = sorted(reports, key=lambda r: (TIER_ORDER.index(r.tier),
                                               -(r.score if r.score is not None else -1.0), r.position))
        for row, rep in enumerate(order, 2):
            values = {"Row #": rep.master_row, "Task ID": rep.task_id, "Agent": rep.agent,
                      "Verdict": verdict_text(rep), "AI Score": rep.score, "Bar": rep.bar,
                      "Decided On": rep.decided_on, "Rationale": rationale_cell(rep), "Trigger": trigger_text(rep),
                      "Each Column Alone": column_scores(rep, self.tier_rules)}
            for name, value in values.items():
                if name in col:
                    ws.cell(row=row, column=col[name], value=value)
            if rep.tier in TIER_FILLS:
                ws.cell(row=row, column=col["Verdict"]).fill = TIER_FILLS[rep.tier]
            for name in ("AI Score", "Bar"):
                ws.cell(row=row, column=col[name]).number_format = "0.0%"
            choices.add(ws.cell(row=row, column=col["Human Reviewer Decision"]))
            _box(ws, row, len(headers), wrap=True)
        ws.freeze_panes = ws.cell(row=2, column=col["Agent"] + 1)
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(order) + 1}"
        widths = {"Row #": 7, "Task ID": 16, "Agent": 28, "Verdict": 24, "AI Score": 9, "Bar": 8, "Decided On": 20,
                  "Rationale": 80, "Trigger": 30, "Each Column Alone": 34, "Human Reviewer Decision": 22,
                  "Reviewer Notes": 30}
        _widths(ws, {col[name]: width for name, width in widths.items() if name in col})

    def add_annotated_master_sheet(self, df: pd.DataFrame, reports: List[RowReport],
                                   quality_results: List[QualityAssessment], row_issues_count: Dict[int, int],
                                   text_cols: Optional[List[str]] = None):
        """The audited rows with a short QA block appended."""
        ws = self.wb.create_sheet(title="Annotated Master")
        orig = list(df.columns)
        present = [c for c in (text_cols or []) if c in df.columns]
        truncated = {i: [c for c in present if not pd.isna(df.loc[i, c]) and len(str(df.loc[i, c])) > EXCEL_CELL_LIMIT]
                     for i in range(len(df))}
        any_issues, any_truncated = any(row_issues_count.values()), any(truncated.values())
        qa = (["QA Verdict", "QA AI Score", "QA Bar", "QA Decided On", "QA Trigger Sentence", "QA Trigger Words",
               "QA Quality Score", "QA Quality Flags"] + (["QA Structural Issues"] if any_issues else [])
              + (["QA Truncated Columns"] if any_truncated else []))
        for col, name in enumerate(orig + qa, 1):
            cell = ws.cell(row=1, column=col, value=name)
            cell.font, cell.fill = FONT_HEADER, NAVY_HEADER if col <= len(orig) else GRAY_HEADER
        for i, rep in enumerate(reports):
            row = i + 2
            for col, name in enumerate(orig, 1):
                value = df.loc[i, name]
                rendered = None if pd.isna(value) else str(value)[:EXCEL_CELL_LIMIT]
                cell = ws.cell(row=row, column=col, value=rendered)
                if isinstance(rendered, str):
                    cell.data_type = "s"
            qual = quality_results[i]
            values = [verdict_text(rep), rep.score, rep.bar, rep.decided_on, trigger_sentences(rep),
                      trigger_words(rep), qual.quality_score, "; ".join(qual.flags) or "None"]
            if any_issues:
                values.append(row_issues_count.get(i, 0))
            if any_truncated:
                values.append(", ".join(truncated[i]) or "None")
            base = len(orig)
            for k, value in enumerate(values, 1):
                ws.cell(row=row, column=base + k, value=value)
            if rep.tier in TIER_FILLS:
                ws.cell(row=row, column=base + 1).fill = TIER_FILLS[rep.tier]
            for k in (2, 3, 7):
                ws.cell(row=row, column=base + k).number_format = "0.0%"
            if any_issues and row_issues_count.get(i, 0):
                ws.cell(row=row, column=base + 9).fill = RED_FILL
        ws.freeze_panes = "A2"
        _autofit_columns(ws, max_len_cap=35)

    def add_technical_sheet(self, reports: List[RowReport], policy):
        """Everything moved out of the other sheets: one line per scored text (each cell, and the row's cells
        checked together)."""
        ws = self.wb.create_sheet(title="Technical")
        headers = ["Row #", "Agent", "Text", "Decision", "Score", "Bar", "Words", "Reason Codes", "Evidence",
                   "Detector Status", "Detector Version", "Policy Version", "Proxy Warnings", "Sentence Drops",
                   "Word Drops"]
        _header(ws, 1, headers)
        row = 1
        for rep in reports:
            texts = dict(rep.cells)
            entries = [(name, res, texts.get(name, "")) for name, res in rep.result.column_results.items()]
            if rep.combined is not None:
                entries.append(("all columns together", rep.combined, JOIN.join(t for _, t in rep.cells if t)))
            for n, (name, res, text) in enumerate(entries):
                row += 1
                values = [rep.master_row, rep.agent, name, res.decision.value, res.score, bar_for(policy, text),
                          policy.word_count(text) if text else 0, ", ".join(res.reason_codes) or "None",
                          "; ".join(e.description for e in res.evidence) or "None", res.detector_status.value,
                          res.detector_version, res.policy_version,
                          "; ".join(rep.proxy_warnings) if n == 0 else "",
                          drops_text(rep.trigger) if name == rep.decided_on else "",
                          word_drops_text(rep.trigger) if name == rep.decided_on else ""]
                for col, value in enumerate(values, 1):
                    ws.cell(row=row, column=col, value=value).font = FONT_DEFAULT
                ws.cell(row=row, column=5).number_format = ws.cell(row=row, column=6).number_format = "0.0%"
                _box(ws, row, len(headers), wrap=True)
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(row, 1)}"
        _widths(ws, {1: 7, 2: 24, 3: 22, 4: 22, 5: 9, 6: 8, 7: 7, 8: 30, 9: 60, 10: 14, 11: 30, 12: 18, 13: 30,
                     14: 60, 15: 40})

    def add_structural_issues_sheet(self, issues: List[StructuralIssue]):
        ws = self.wb.create_sheet(title="Structural Issues")
        _header(ws, 1, ["Excel Row", "Column", "Severity", "Rule Violated", "Offending Value", "Description"])
        for row, issue in enumerate(issues, start=2):
            ws.cell(row=row, column=1, value=issue.excel_row).alignment = Alignment(horizontal="center")
            ws.cell(row=row, column=2, value=issue.column)
            severity = ws.cell(row=row, column=3, value=issue.severity)
            severity.alignment = Alignment(horizontal="center")
            severity.fill = RED_FILL if issue.severity == "ERROR" else YELLOW_FILL
            if issue.severity == "ERROR":
                severity.font = FONT_BOLD
            ws.cell(row=row, column=4, value=issue.rule_violated)
            ws.cell(row=row, column=5, value=str(issue.value)[:80])
            ws.cell(row=row, column=6, value=issue.description)
            _box(ws, row, 6)
        _autofit_columns(ws, max_len_cap=55)

    def save(self, filepath: str):
        # Reports contain no executable formulas. Keep every formula-like value
        # originating from uploaded data or detector evidence as inert text.
        for worksheet in self.wb.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        cell.data_type = "s"
        self.wb.save(filepath)
