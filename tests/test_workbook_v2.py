"""The v2 audit workbook (spec 2026-10-01, section 4)."""
import re
import zipfile
from types import SimpleNamespace

import openpyxl
import pandas as pd
from openpyxl.cell.rich_text import CellRichText, TextBlock

from src.config import PipelineConfig
from src.pipeline import QAPipeline
from src.report.excel_writer import rationale_cell
from src.report.triggers import Phrase, TriggerResult
from src.report.tiers import HIGH, LIKELY_AI, LOW, NOT_CHECKED
from tests.test_row_reports import PLANT, _service

ROWS = [
    {"rater_id": "ann", "task_id": "T1", "why": f"The rating is fair overall. {PLANT} It matches the rubric well."},
    {"rater_id": "bo", "task_id": "T2", "why": "The rating is fair and the response cites the source it used."},
    {"rater_id": "ann", "task_id": "T3", "why": "It is certainly a clear and complete answer to the question."},
    {"rater_id": "cy", "task_id": "T4", "why": "Too short."},
]


def _audit(tmp_path, rows=ROWS, cols=("why",)):
    src, out = tmp_path / "in.csv", tmp_path / "out.xlsx"
    pd.DataFrame(rows).to_csv(src, index=False)
    config = PipelineConfig()
    config.structural.text_columns = list(cols)
    config.structural.skip_unassigned = False
    QAPipeline(config=config, authorship_service=_service()).run(str(src), str(out))
    return out


def _headers(ws):
    return [c.value for c in ws[1]]


def test_sheets_in_order_and_no_structural_sheet_without_issues(tmp_path):
    wb = openpyxl.load_workbook(_audit(tmp_path, ROWS[:2]))   # ann twice would be a duplicate warning
    assert wb.sheetnames == ["Summary Dashboard", "Agents", "Trigger Words", "AI Review", "Annotated Master",
                             "Technical"]


def test_ai_review_has_one_verdict_the_bar_and_the_full_rationale_strongest_first(tmp_path):
    ws = openpyxl.load_workbook(_audit(tmp_path))["AI Review"]
    h = _headers(ws)
    assert h == ["Row #", "Task ID", "Agent", "Verdict", "AI Score", "Bar", "Decided On", "Rationale", "Trigger",
                 "Each Column Alone", "Human Reviewer Decision", "Reviewer Notes"]
    col = {name: i + 1 for i, name in enumerate(h)}
    verdicts = [ws.cell(r, col["Verdict"]).value for r in range(2, 6)]
    assert verdicts == [LIKELY_AI, HIGH, LOW, f"{NOT_CHECKED} (too short)"]
    assert ws.cell(3, col["Bar"]).value == 0.99 and ws.cell(3, col["Task ID"]).value == "T1"
    assert ws.cell(3, col["Rationale"]).value == ROWS[0]["why"]          # full text, not cut
    assert ws.cell(3, col["Trigger"]).value == ('Without it: 95% -> 20% (100% of the AI signal)\n'
                                                'Words: "rich tapestry" (100%)')
    assert "Review Tag" not in h and "Detector Version" not in h


def test_the_trigger_sentence_is_bold_red_and_its_words_underlined_inside_the_rationale(tmp_path):
    ws = openpyxl.load_workbook(_audit(tmp_path), rich_text=True)["AI Review"]
    col = _headers(ws).index("Rationale") + 1
    cell = ws.cell(3, col).value
    assert isinstance(cell, CellRichText)
    red = [b for b in cell if isinstance(b, TextBlock) and b.font.b]
    assert "".join(b.text for b in red) == PLANT
    assert all(str(b.font.color.rgb).endswith("C00000") for b in red)
    assert [b.text for b in red if b.font.u] == ["rich tapestry"]


def test_trigger_words_outside_a_trigger_sentence_are_underlined_in_orange(tmp_path):
    ws = openpyxl.load_workbook(_audit(tmp_path), rich_text=True)["AI Review"]
    col = _headers(ws).index("Rationale") + 1
    cell = ws.cell(2, col).value                                   # the single-sentence "certainly" row
    marked = [b for b in cell if isinstance(b, TextBlock)]
    assert [b.text for b in marked] == ["certainly"]
    assert marked[0].font.u and str(marked[0].font.color.rgb).endswith("C65911")


def test_the_trigger_words_sheet_lists_each_word_with_how_many_rows_contain_it(tmp_path):
    ws = openpyxl.load_workbook(_audit(tmp_path))["Trigger Words"]
    assert _headers(ws) == ["Word or Phrase", "Rows It Triggered", "Agents", "Average Share of AI Signal",
                            "Rows Containing It", "Share of All Rows", "Agent Names"]
    rows = {ws.cell(r, 1).value: [ws.cell(r, c).value for c in range(2, 8)] for r in range(2, ws.max_row + 1)}
    assert rows["rich tapestry"] == [1, 1, 1.0, 1, 0.25, "ann"]


def test_agents_sheet_ranks_agents(tmp_path):
    ws = openpyxl.load_workbook(_audit(tmp_path))["Agents"]
    assert _headers(ws)[:4] == ["Agent", "Rows", LIKELY_AI, HIGH]
    assert [ws.cell(r, 1).value for r in range(2, 5)] == ["ann", "bo", "cy"]


def test_master_has_the_short_qa_block(tmp_path):
    ws = openpyxl.load_workbook(_audit(tmp_path))["Annotated Master"]
    qa = [h for h in _headers(ws) if str(h).startswith("QA ")]
    assert qa[:8] == ["QA Verdict", "QA AI Score", "QA Bar", "QA Decided On", "QA Trigger Sentence",
                      "QA Trigger Words", "QA Quality Score", "QA Quality Flags"]
    col = _headers(ws).index("QA Trigger Sentence") + 1
    assert ws.cell(2, col).value == PLANT
    assert ws.cell(2, col + 1).value == "rich tapestry (100%)"
    assert ws.cell(4, col).value in (None, "") and ws.cell(4, col + 1).value == "certainly (100%)"


def test_technical_sheet_has_a_line_per_scored_text_with_sentence_drops(tmp_path):
    rows = [{"rater_id": "a", "first": "The rating is fair overall here.", "second": f"{PLANT} It matches."}]
    ws = openpyxl.load_workbook(_audit(tmp_path, rows, cols=("first", "second")))["Technical"]
    texts = [ws.cell(r, 3).value for r in range(2, ws.max_row + 1)]
    assert texts == ["first", "second", "all columns together"]
    drops_col = _headers(ws).index("Sentence Drops") + 1
    assert "of signal" in ws.cell(4, drops_col).value and ws.cell(2, drops_col).value in (None, "")
    words_col = _headers(ws).index("Word Drops") + 1
    assert '"rich"' in ws.cell(4, words_col).value and ws.cell(2, words_col).value in (None, "")


def test_summary_has_one_tier_table_with_rates_or_a_not_measured_note(tmp_path):
    ws = openpyxl.load_workbook(_audit(tmp_path))["Summary Dashboard"]
    text = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert HIGH in text and "Agents with any Likely AI" in text and "Trigger sentences and words" in text
    assert any(t.startswith("Trigger Words:") for t in text)
    assert "not measured for this model" in text                       # planted-v1 has no measured rates
    assert "Likely Human-written" not in text and "Review Tags (strongest first)" not in text


def test_a_formula_like_rationale_with_a_trigger_stays_text(tmp_path):
    rows = [{"rater_id": "a", "why": f"=SUM(1,2) is what the rater typed. {PLANT} It ends here."}]
    wb = openpyxl.load_workbook(_audit(tmp_path, rows))
    assert all(c.data_type != "f" for ws in wb.worksheets for row in ws.iter_rows() for c in row)


def test_the_summary_says_when_no_agent_column_was_found(tmp_path):
    rows = [{"prompt_id": "P1", "why": "The rating is fair and the response cites the source it used."}]
    ws = openpyxl.load_workbook(_audit(tmp_path, rows))["Summary Dashboard"]
    facts = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(6, 16)}
    assert facts["Agent column"] == "none found"


def test_a_rationale_never_writes_a_whitespace_only_text_run(tmp_path):
    """openpyxl writes a whitespace-only run without xml:space="preserve", which Excel repairs as broken "String
    properties" on opening (the space between two adjacent trigger words, or the blank line between two cells)."""
    first, second = "It keeps some nuance overall.", "Fine."
    words = [Phrase("a", 9, 13, "some", 0.5), Phrase("a", 14, 20, "nuance", 0.4)]
    report = SimpleNamespace(trigger=TriggerResult(kind="spread", words=words),
                             decided=[("a", first), ("b", second)])
    out = tmp_path / "rich.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = rationale_cell(report)
    wb.save(out)
    xml = zipfile.ZipFile(out).read("xl/worksheets/sheet1.xml").decode("utf8")
    assert re.findall(r"<t>\s+</t>", xml) == []
    back = openpyxl.load_workbook(out, rich_text=True).active["A1"].value
    assert str(back).replace("\r\n", "\n") == f"[a] {first}\n\n[b] {second}"
    assert [b.text.strip() for b in back if isinstance(b, TextBlock) and b.font.u] == ["some", "nuance"]
