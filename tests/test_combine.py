"""Combined multi-locale workbook (src/report/combine.py)."""
import openpyxl
import pytest
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

from src.report.combine import CombineError, combine_workbooks

HIGH = "High AI score – review"
RED = InlineFont(b=True, color="C00000")


def _audit(path, locale, verdicts, extra=0):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary Dashboard"
    for row in (["EVALUATION QA AUDIT REPORT"], [], ["Run", "Value"], ["Model", "m-1"], ["Locale", locale],
                ["Structural errors", 0], [], ["Tier", "Rows", "Meaning", "H25", "A25", "Hall", "Aall"],
                ["Likely AI", 1, "bar", 0.001, 0.3, 0.002, 0.4], [HIGH, 1, "review", 0.06, 0.5, 0.07, 0.5],
                ["Not checked", 1, "none"], ["A note under the tiers."]):
        ws.append(row)
    wb.create_sheet("Agents").append(["Agent", "Rows", "Likely AI"])
    wb["Agents"].append(["a@x.ai", len(verdicts), verdicts.count("Likely AI")])
    wb.create_sheet("Trigger Words").append(["Word or Phrase", "Rows It Triggered"])
    review = wb.create_sheet("AI Review")
    review.append(["Row #", "Task ID", "Agent", "Verdict", "AI Score", "Bar", "Decided On", "Rationale",
                   "Human Reviewer Decision"])
    review.append([2, "t1", "a@x.ai", verdicts[0], 0.99, 0.98, "col", CellRichText(TextBlock(RED, "red sentence"), " rest"), None])
    master = wb.create_sheet("Annotated Master")
    pad = [None] * extra  # each locale's own input columns push QA Verdict to a different position
    master.append(["Task ID", "Agent", "locale", "replicate", "skip", "comment"] + ["input"] * extra + ["QA Verdict"])
    for i, verdict in enumerate(verdicts):
        master.append([f"t{i}", "a@x.ai", locale, 1, None, "text"] + pad + [verdict])
    wb.save(path)


def test_combines_locales_into_tabs(tmp_path):
    a, b, out = tmp_path / "audited_aa_x_en-US.xlsx", tmp_path / "audited_bb_x_fr-FR.xlsx", tmp_path / "all.xlsx"
    _audit(a, "en-US", ["Likely AI", HIGH, "Not checked (too short)"])
    _audit(b, "fr-FR", [HIGH, HIGH], extra=6)
    result = combine_workbooks([str(a), str(b)], str(out))
    assert result["rows"] == 5 and result["locales"] == ["en-US", "fr-FR"]

    wb = openpyxl.load_workbook(out, rich_text=True)
    assert wb.sheetnames == ["All Locales", "en-US", "fr-FR", "Trigger Words", "Agents (all)", "Annotated Master"]
    overview = {r[0]: r for r in wb["All Locales"].iter_rows(values_only=True) if r[0]}
    assert overview["Rows"][1] == 5 and overview["en-US"][1] == 3 and overview["fr-FR"][1] == 2
    assert overview["Likely AI"][1] == 1 and overview[HIGH][1] == 3 and overview["Not checked"][1] == 1
    cell = wb["en-US"]["H2"].value  # the red trigger run survives
    assert isinstance(cell, CellRichText) and cell[0].font.color.rgb.endswith("C00000")
    assert [r[0] for r in wb["Agents (all)"].iter_rows(min_row=2, values_only=True)] == ["en-US", "fr-FR"]


def test_duplicate_locale_names_get_distinct_tabs(tmp_path):
    paths = []
    for name in ("one", "two"):
        paths.append(str(tmp_path / f"{name}.xlsx"))
        _audit(paths[-1], "en-US", ["Likely AI"])
    out = tmp_path / "all.xlsx"
    assert combine_workbooks(paths, str(out))["locales"] == ["en-US", "en-US-2"]


def test_needs_two_audit_workbooks(tmp_path):
    a = tmp_path / "a.xlsx"
    _audit(a, "en-US", ["Likely AI"])
    with pytest.raises(CombineError):
        combine_workbooks([str(a)], str(tmp_path / "o.xlsx"))
    plain = tmp_path / "plain.xlsx"
    openpyxl.Workbook().save(plain)
    with pytest.raises(CombineError, match="not an audit workbook"):
        combine_workbooks([str(a), str(plain)], str(tmp_path / "o.xlsx"))
