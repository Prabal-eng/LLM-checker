"""End-to-end integration test for the pipeline."""
import os
import openpyxl
import pandas as pd
from src.pipeline import QAPipeline
from src.config import PipelineConfig


def test_pipeline_end_to_end(tmp_path):
    # Create test input file
    input_file = tmp_path / "test_input.xlsx"
    output_file = tmp_path / "test_output.xlsx"

    df = pd.DataFrame([
        {
            "rater_id": "RATER_1",
            "row_id": "T1",
            "category": "Classification",
            "verdict": "PASS",
            "justification": "Checked against reference dataset, everything lines up properly."
        },
        {
            "rater_id": "RATER_2",
            "row_id": "T2",
            "category": "Reasoning",
            "verdict": "PASS",
            "justification": "Furthermore, it is crucial to delve into the comprehensive nature of this response. In conclusion, it serves as a testament to perfection."
        },
        {
            "rater_id": "RATER_3",
            "row_id": "T3",
            "category": "Math",
            "verdict": "FAIL",
            "justification": ""
        }
    ])
    df.to_excel(input_file, index=False)

    config = PipelineConfig()
    pipeline = QAPipeline(config=config)

    res = pipeline.run(str(input_file), str(output_file))

    assert res.total_rows == 3
    assert res.ai_flagged_count == 0
    assert res.likely_human_count == 0
    assert res.inconclusive_count == 3
    assert res.structural_errors_count >= 1
    assert os.path.exists(str(output_file))

    # Verify generated sheets
    wb = openpyxl.load_workbook(str(output_file))
    assert "Summary Dashboard" in wb.sheetnames
    assert "AI Review" in wb.sheetnames
    assert "Structural Issues" in wb.sheetnames
    assert "Annotated Master" in wb.sheetnames
    review_headers = [cell.value for cell in wb["AI Review"][1]]
    assert "Verdict" in review_headers and "Bar" in review_headers
    assert "Technical" in wb.sheetnames and "Agents" in wb.sheetnames


def test_csv_text_over_excel_limit_is_explicitly_rejected_and_marked(tmp_path):
    input_file = tmp_path / "long.csv"
    output_file = tmp_path / "long_audited.xlsx"
    original = "word " * 8_000
    pd.DataFrame([{
        "rater_id": "RATER_1",
        "row_id": "T1",
        "verdict": "PASS",
        "justification": original,
    }]).to_csv(input_file, index=False)

    result = QAPipeline().run(str(input_file), str(output_file))

    assert result.inconclusive_count == 1
    wb = openpyxl.load_workbook(output_file)
    ws = wb["Annotated Master"]
    headers = [cell.value for cell in ws[1]]
    values = {headers[index]: ws.cell(row=2, column=index + 1).value for index in range(len(headers))}
    assert len(values["justification"]) == 32_767
    assert values["QA Verdict"] == "Not checked (too long)"
    assert values["QA Truncated Columns"] == "justification"


def test_csv_formula_like_text_remains_inert_in_workbook(tmp_path):
    input_file = tmp_path / "formula.csv"
    output_file = tmp_path / "formula_audited.xlsx"
    formula_text = '=HYPERLINK("http://example.com","click")'
    pd.DataFrame([{
        "rater_id": "RATER_1",
        "row_id": "T1",
        "verdict": "PASS",
        "justification": formula_text,
    }]).to_csv(input_file, index=False)

    QAPipeline().run(str(input_file), str(output_file))

    workbook = openpyxl.load_workbook(output_file, data_only=False)
    ws = workbook["Annotated Master"]
    headers = [cell.value for cell in ws[1]]
    cell = ws.cell(row=2, column=headers.index("justification") + 1)
    assert cell.value == formula_text
    assert cell.data_type == "s"
    assert all(
        cell.data_type != "f"
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
    )


def _counting_service():
    from src.detection import AuthorshipService, TextContext

    class Counting(AuthorshipService):
        def check_text(self, text, context=TextContext.GENERAL):
            self.text_calls += 1
            return super().check_text(text, context)

        def check_batch(self, texts, context=TextContext.GENERAL):
            self.batches.append(list(texts))
            return super().check_batch(texts, context)

    svc = Counting(QAPipeline(PipelineConfig()).detector, model_validated=False, detector_version="t")
    svc.text_calls, svc.batches = 0, []
    return svc


def test_prefetch_scores_each_cell_and_each_joined_row_once_in_batches():
    svc = _counting_service()
    pipeline = QAPipeline(PipelineConfig(), authorship_service=svc)
    df = pd.DataFrame({"a": ["x one", "same", None], "b": ["y two", "same", "z"]})
    pipeline.prefetch_sheet(df, ["a", "b"])
    scored = [t for batch in svc.batches for t in batch]
    assert sorted(scored) == sorted(["x one", "y two", "x one\n\ny two", "same", "same\n\nsame", "", "z"])
    assert svc.text_calls == 0
    pipeline._cached_detect("x one")
    assert svc.text_calls == 0 and len(svc.batches) == 1  # served from the prefetched cache


def test_a_sheet_run_scores_in_batches_not_cell_by_cell(tmp_path):
    input_file, output_file = tmp_path / "in.xlsx", tmp_path / "out.xlsx"
    pd.DataFrame([{"rater_id": f"R{i}", "row_id": f"T{i}", "category": "Math", "verdict": "PASS",
                   "justification": f"The steps in answer {i} check out against the reference."} for i in range(5)]
                 ).to_excel(input_file, index=False)
    svc = _counting_service()
    QAPipeline(PipelineConfig(), authorship_service=svc).run(str(input_file), str(output_file))
    assert svc.text_calls == 0 and len(svc.batches) == 1
