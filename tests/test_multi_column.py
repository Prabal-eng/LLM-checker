"""Unit tests for multi-column and custom column selection for AI detection."""
import os
import pandas as pd
import openpyxl
from src.pipeline import QAPipeline
from src.config import PipelineConfig
from src.detection import AuthorshipService
from src.detectors.base import BaseDetector, DetectionResult


class KeywordDetector(BaseDetector):
    def detect(self, text: str) -> DetectionResult:
        score = 0.95 if "Furthermore" in text else 0.10
        return DetectionResult(
            provider="test",
            ai_score=score,
            is_flagged=score >= 0.80,
            analysis={"evidence": ["deterministic test signal"]},
        )


def test_multi_column_detection(tmp_path):
    input_file = tmp_path / "multi_col_test.xlsx"
    output_file = tmp_path / "multi_col_audited.xlsx"

    # Dataset with 2 separate rationale columns
    data = [
        {
            "rater_id": "alex.chen",
            "task_id": "T1",
            "category": "Classification",
            "verdict": "PASS",
            "image_factuality_rationale": "The generated image accurately reproduces the architectural features.",
            "overall_assessment_rationale": "Furthermore, it is crucial to delve into the comprehensive nature of this synthesis. In conclusion, it serves as a testament to perfection."
        },
        {
            "rater_id": "bob.smith",
            "task_id": "T2",
            "category": "Reasoning",
            "verdict": "PASS",
            "image_factuality_rationale": "Clear evidence aligns closely with every requested prompt detail here.",
            "overall_assessment_rationale": "Solid performance meets all stated criteria with specific supporting details."
        },
        {
            "rater_id": "casey.jones",
            "task_id": "T3",
            "category": "Reasoning",
            "verdict": "PASS",
            "image_factuality_rationale": "Too short",
            "overall_assessment_rationale": "Also short"
        }
    ]
    pd.DataFrame(data).to_excel(input_file, index=False)

    config = PipelineConfig()
    config.structural.text_columns = ["image_factuality_rationale", "overall_assessment_rationale"]
    service = AuthorshipService(
        KeywordDetector(), model_validated=True, detector_version="test-v1"
    )
    pipeline = QAPipeline(config=config, authorship_service=service)

    res = pipeline.run(str(input_file), str(output_file))

    assert res.total_rows == 3
    # Row 1 should be flagged because overall_assessment_rationale has AI buzzwords
    assert res.ai_flagged_count >= 1
    # Spreadsheet text is rater-comment context, which is flag-only: missed AI is unmeasured
    # there, so a clean row is Inconclusive rather than Likely human.
    assert res.likely_human_count == 0
    assert res.inconclusive_count == 2
    assert os.path.exists(str(output_file))

    wb = openpyxl.load_workbook(str(output_file))
    assert "AI Review" in wb.sheetnames
    assert "Annotated Master" in wb.sheetnames

    master_ws = wb["Annotated Master"]
    header_vals = [cell.value for cell in master_ws[1]]
    assert "QA Verdict" in header_vals and "QA Decided On" in header_vals
    assert any("image_factuality" in str(h) for h in header_vals)
    assert any("overall_assessment" in str(h) for h in header_vals)
    assert not any(str(h).startswith("QA_Text_") for h in header_vals)   # per-cell detail is on Technical

    technical = wb["Technical"]
    texts = [technical.cell(row=r, column=3).value for r in range(2, technical.max_row + 1)]
    assert texts.count("image_factuality_rationale") == 3 and texts.count("all columns together") == 3
    tech_headers = [cell.value for cell in technical[1]]
    assert "Detector Version" in tech_headers and "Policy Version" in tech_headers

    review_headers = [cell.value for cell in wb["AI Review"][1]]
    assert "Each Column Alone" in review_headers and "Decided On" in review_headers
